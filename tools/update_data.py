#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Автообновление данных сайта FIRST® Kazakhstan.

Запускается GitHub Action'ом (.github/workflows/update-site.yml) по расписанию, вручную и при каждом изменении
репозитория. Можно запустить и на компьютере:  pip install openpyxl  &&  python tools/update_data.py
Проверить код без сети (то же делает Actions для каждого pull request):  python tools/selfcheck.py

1. Скачивает официальную Google-таблицу расписания (файл .xlsx): значения и ссылки всех вкладок
   -> assets/data/sheet.js. Сайт сам собирает из неё карточки событий (даты, места, категории, регистрация).
2. Открывает каждую форму регистрации из таблицы: закрытая форма или форма другого события
   (например, прошлогодняя) помечается, и сайт не показывает для неё кнопку.
3. Из той же таблицы читает необязательные вкладки «Сайт» (название и тема сезона, логотип, видео, ссылки на правила,
   письмо и расписание, подтемы FLL/FTC/FRC) и «Новости» (новости сайта) -> assets/data/site.js,
   логотип сезона -> assets/img/season.png|jpg|webp, фото новостей -> assets/img/news/sh-*.
4. Берёт посты публичного Telegram-канала с хэштегом из tools/config.json (по умолчанию #сайт); telegram_max_news = 0 – хранить все
   -> assets/data/news.js, фото -> assets/img/news/.
5. Открывает список событий uSTEM Connect (connect.ustem.kz/events — все события, и отдельно только с открытой регистрацией):
   адреса их страниц -> assets/data/sheet.js (connect_all, connect). Сайт сам находит в них страницу турнира по названию и ведёт кнопку
   регистрации прямо туда; пока регистрация открыта, над кнопкой пометка «Регистрация открыта в uSTEM Connect».
6. Время последней успешной проверки -> assets/data/status.js (служебное; на сайте «Обновлено» = дата изменения таблицы из sheet.js).
Если какой-то источник недоступен, прежние данные остаются как есть: сайт никогда не остаётся пустым.
"""
import datetime, hashlib, html, io, json, os, re, sys, time, urllib.error, urllib.parse, urllib.request

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, 'assets', 'data')
IMG = os.path.join(ROOT, 'assets', 'img')
NEWS_IMG = os.path.join(ROOT, 'assets', 'img', 'news')
CFG = json.load(open(os.path.join(ROOT, 'tools', 'config.json'), encoding='utf-8'))
UA = 'Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124 Safari/537.36 FIRST-Kazakhstan-site-updater'
YEAR = re.compile(r'20\d\d')
STAGES = ('scrimmage', 'regional', 'qualifier', 'kick-off', 'championship', 'off-season')
# Официальные сайты, на которые могут вести ссылки из вкладки «Сайт» (правила, письмо, расписание); «www.» впереди допускается.
LINK_HOSTS = ('docs.google.com', 'drive.google.com', 'firstinspires.org', 'firstlegoleague.org', 'first.global', 'firstrobotics.kz', 'ustem.kz')
VIDEO_HOSTS = ('youtube.com', 'youtu.be')
# Скрипт ходит только по адресам Google, Telegram и официальных сайтов FIRST (и по их переадресациям); всё остальное из таблицы не открывается.
CONNECT_HOST = 'connect.ustem.kz'   # uSTEM Connect — регистрация на официальные турниры (как CONNECT в assets/js/data.js)
ALLOWED_HOSTS = ('forms.gle', 'docs.google.com', 't.me', CONNECT_HOST) + LINK_HOSTS + tuple('www.' + h for h in LINK_HOSTS[2:])
ALLOWED_SUFFIXES = ('.google.com', '.googleusercontent.com', '.telegram.org', '.telesco.pe', '.cdn-telegram.org', '.t.me')


def host_ok(url):
    try:
        p = urllib.parse.urlsplit(url)
    except ValueError:
        return False
    h = (p.hostname or '').lower()
    return p.scheme == 'https' and bool(h) and (h in ALLOWED_HOSTS or h.endswith(ALLOWED_SUFFIXES))


def link_ok(url, hosts=LINK_HOSTS):
    """https-ссылка на один из разрешённых сайтов («www.» впереди не считается)."""
    try:
        p = urllib.parse.urlsplit(url or '')
    except ValueError:
        return False
    h = (p.hostname or '').lower()
    return p.scheme == 'https' and (h[4:] if h.startswith('www.') else h) in hosts


def https_ok(url):
    """Любая https-ссылка (javascript:, data:, http: и прочее не проходят)."""
    try:
        p = urllib.parse.urlsplit(url or '')
    except ValueError:
        return False
    return p.scheme == 'https' and bool(p.hostname)


FORM_RX = re.compile(r'^https://(forms\.gle/|docs\.google\.com/forms/)')   # для журналов и тестов; проверка — is_form_url()


def is_form_url(url):
    """Ссылка на Google Форму: только https, только forms.gle или docs.google.com/forms/…"""
    try:
        p = urllib.parse.urlsplit(url or '')
    except ValueError:
        return False
    h = (p.hostname or '').lower()
    return p.scheme == 'https' and (h == 'forms.gle' or (h == 'docs.google.com' and p.path.startswith('/forms/')))


class _Redirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        if not host_ok(newurl):
            raise urllib.error.HTTPError(newurl, code, 'переадресация на недоверенный адрес: ' + newurl[:80], headers, fp)
        return super().redirect_request(req, fp, code, msg, headers, newurl)


OPENER = urllib.request.build_opener(_Redirect)


def get(url, timeout=40, cap=20 * 1024 * 1024):
    if not host_ok(url):
        raise ValueError('недоверенный адрес: ' + str(url)[:80])
    req = urllib.request.Request(url, headers={'User-Agent': UA, 'Accept-Language': 'ru,kk;q=0.9,en;q=0.8'})
    with OPENER.open(req, timeout=timeout) as r:
        data = r.read(cap + 1)
        if len(data) > cap:
            raise RuntimeError('ответ больше %d МБ: %s' % (cap // (1024 * 1024), url[:80]))
        return r.geturl(), data


def image_kind(data):
    """Расширение по первым байтам файла: png / jpg / webp, иначе None."""
    if data.startswith(b'\x89PNG'):
        return 'png'
    if data.startswith(b'\xff\xd8\xff'):
        return 'jpg'
    if data[:4] == b'RIFF' and data[8:12] == b'WEBP':
        return 'webp'
    return None


def drive_direct(url):
    """Ссылка «Поделиться» Google Диска (…/file/d/ID/view, …/open?id=ID) -> прямая ссылка на скачивание файла."""
    m = re.match(r'https://drive\.google\.com/(?:file/d/|open\?id=|uc\?(?:export=\w+&)?id=)([\w-]+)', url or '')
    return 'https://drive.google.com/uc?export=download&id=' + m.group(1) if m else url


def fetch_image(url, cap=3 * 1024 * 1024):
    """Скачать картинку (PNG/JPEG/WebP не больше cap) с доверенного https-адреса -> (байты, расширение)."""
    url = drive_direct((url or '').strip())
    if not host_ok(url):
        raise ValueError('картинки берутся только с Google Диска, Telegram и официальных сайтов FIRST (https)')
    _, data = get(url, 30, cap=cap)
    kind = image_kind(data)
    if not kind:
        raise RuntimeError('это не картинка (нужен PNG, JPEG или WebP; ссылка должна вести на сам файл)')
    return data, kind


def now_iso():
    return datetime.datetime.now(datetime.timezone.utc).replace(microsecond=0).isoformat().replace('+00:00', 'Z')


def read_js(name, var):
    """Прочитать JSON из assets/data/<name> (формат: window.VAR=<json>;)"""
    p = os.path.join(DATA, name)
    if not os.path.exists(p):
        return None
    s = open(p, encoding='utf-8').read()
    i = s.find('window.' + var + '=')
    if i < 0:
        return None
    s = s[i + len('window.' + var + '='):].strip()
    if s.endswith(';'):
        s = s[:-1]
    try:
        return json.loads(s)
    except ValueError:
        return None


def write_js(name, var, obj, comment):
    p = os.path.join(DATA, name)
    js = json.dumps(obj, ensure_ascii=False, separators=(',', ':'))
    js = js.replace('\u2028', '\\u2028').replace('\u2029', '\\u2029').replace('<', '\\u003c')   # безопасно для старых движков и при вставке в HTML
    body = '/* ' + comment + ' */\nwindow.' + var + '=' + js + ';\n'
    tmp = p + '.tmp'
    open(tmp, 'w', encoding='utf-8').write(body)
    os.replace(tmp, p)


# ---------------------------------------------------------------- 1. таблица
def cell_text(v):
    if v is None:
        return ''
    if isinstance(v, datetime.datetime):
        return v.strftime('%d.%m.%Y') if (v.hour, v.minute, v.second) == (0, 0, 0) else v.strftime('%d.%m.%Y %H:%M')
    if isinstance(v, datetime.date):
        return v.strftime('%d.%m.%Y')
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def fetch_sheet(sheet_id):
    import openpyxl
    _, data = get('https://docs.google.com/spreadsheets/d/%s/export?format=xlsx' % sheet_id, 120)
    if not data.startswith(b'PK'):
        raise RuntimeError('таблица не открылась как .xlsx (нет доступа по ссылке?)')
    wv = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    wf = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    tabs = {}
    for ws in wv.worksheets:
        if ws.sheet_state != 'visible' and not special_tab(ws.title):   # скрытые вкладки событий (старые сезоны) не читаем; «Сайт»/«Новости» можно скрыть
            continue
        ff = wf[ws.title]
        rows = []
        for r in range(1, min(ws.max_row, 500) + 1):
            row = []
            for c in range(1, min(ws.max_column, 40) + 1):
                cell = ws.cell(r, c)
                t = cell_text(cell.value)
                url = cell.hyperlink.target if (cell.hyperlink is not None and cell.hyperlink.target) else ''
                if not url:
                    f = ff.cell(r, c).value
                    if isinstance(f, str) and f.upper().startswith('=HYPERLINK('):
                        m = re.search(r'HYPERLINK\(\s*"([^"]+)"(?:\s*[;,]\s*"([^"]*)")?', f, re.I)
                        if m:
                            url = m.group(1)
                            if not t and m.group(2):
                                t = m.group(2)
                row.append([t, url] if url else t)
            while row and row[-1] == '':
                row.pop()
            rows.append(row)
        while rows and not rows[-1]:
            rows.pop()
        if rows:
            tabs[ws.title] = rows
    return tabs


def special_tab(title):
    """'site' для вкладки «Сайт», 'news' для «Новости» (регистр и пробелы не важны), иначе None — это вкладка событий."""
    t = (title or '').strip().lower()
    return 'site' if t in ('сайт', 'site') else 'news' if t in ('новости', 'news', 'жаңалықтар') else None


def txt(c):
    return c[0] if isinstance(c, list) else (c or '')


def cell_url(c):
    """Адрес из ячейки: цель гиперссылки или сам текст ячейки (проверяется дальше через link_ok/https_ok)."""
    return (c[1] if isinstance(c, list) else (c or '')).strip()


def name_col(head):
    for i, h in enumerate(head):
        if re.search(r'атау|название|name', txt(h).lower()):
            return i
    return -1


def date_col(head):
    for i, h in enumerate(head):
        if re.search(r'күн|дата|date', txt(h).lower()):
            return i
    return -1


def count_events(tabs):
    n = 0
    for rows in tabs.values():
        nc = name_col(rows[0]) if rows else -1
        n += sum(1 for r in rows[1:] if 0 <= nc < len(r) and txt(r[nc]).strip())
    return n


def validate(tabs, prev_tabs):
    good = [t for t, rows in tabs.items() if rows and name_col(rows[0]) >= 0 and date_col(rows[0]) >= 0]
    if not good:
        raise RuntimeError('в таблице не найдены вкладки с колонками «Атауы» и «Күні»')
    n, p = count_events(tabs), count_events(prev_tabs or {})
    if p >= 20 and n < p * 0.3:
        raise RuntimeError('в таблице %d событий вместо %d — похоже на случайное удаление, данные не обновлены' % (n, p))
    return n


# ---------------------------------------------------------------- 2. формы регистрации
def form_events(tabs):
    """url формы -> список событий (название, дата) из той же строки"""
    out = {}
    for rows in tabs.values():
        if not rows:
            continue
        nc, dc = name_col(rows[0]), date_col(rows[0])
        cur = ''
        for r in rows[1:]:
            name = txt(r[nc]).strip() if 0 <= nc < len(r) else ''
            if name:
                cur = name
            date = txt(r[dc]) if 0 <= dc < len(r) else ''
            for c in r:
                if isinstance(c, list) and is_form_url(c[1]):
                    out.setdefault(c[1], []).append({'name': cur, 'date': date})
    return out


# слова, которые встречаются в названиях многих событий и не говорят о «чужой» форме
GENERIC = {'kazakhstan', 'kyrgyzstan', 'uzbekistan', 'qazaqstan', 'central', 'asia', 'online', 'national', 'republic'}


def brand(name):
    m = re.match(r'\s*([A-Za-z][A-Za-z\-]+)', name or '')
    w = m.group(1).lower() if m else ''
    return '' if w == 'first' else w


def mismatch(title, ev, brands):
    """True, если заголовок формы явно относится к другому событию."""
    t, n = (title or '').lower(), ev['name'].lower()
    if not t:
        return False
    yt, yn = set(YEAR.findall(t)), set(YEAR.findall(n))
    if yt and yn and not (yt & yn):
        return True                                   # другой сезон / год
    st, sn = {s for s in STAGES if s in t}, {s for s in STAGES if s in n}
    if st and sn and not (st & sn):
        return True                                   # другой этап (например, форма скримиджа у регионального этапа)
    own = brand(ev['name'])
    if own and own in t:
        return False
    others = [b for b in brands if b and b != own and b not in GENERIC and re.search(r'(?<![a-z])' + re.escape(b) + r'(?![a-z])', t)]
    return bool(others)                               # в заголовке другой город/событие


def page_title(page):
    m = re.search(r'<meta[^>]+property="og:title"[^>]+content="([^"]*)"', page) or re.search(r'<title>(.*?)</title>', page, re.S)
    return html.unescape(m.group(1)).strip() if m else ''


def check_forms(tabs, prev):
    links = form_events(tabs)
    brands = {brand(e['name']) for evs in links.values() for e in evs}
    for rows in tabs.values():
        nc = name_col(rows[0]) if rows else -1
        brands |= {brand(txt(r[nc])) for r in rows[1:] if 0 <= nc < len(r)}
    brands.discard('')
    out = {}
    for url, evs in links.items():
        if re.search(r'/forms/d/[^/]+/edit', url):      # ссылка на редактор формы, посетитель её не откроет
            out[url] = {'s': 'mismatch', 't': 'ссылка ведёт на редактор формы, а не на саму форму'}
            print('  [редактор формы] %s — %s' % (url, ', '.join(sorted({e['name'] for e in evs}))))
            continue
        try:
            final, body = get(url, 15)
            page = body.decode('utf-8', 'replace')
            title = page_title(page)
            closed = '/closedform' in final or bool(re.search(
                r'no longer accepting responses|больше не принимает ответы|жауаптарды қабылдамайды', page, re.I))
            if closed:
                st = 'closed'
            elif all(mismatch(title, e, brands) for e in evs):
                st = 'mismatch'
            else:
                st = 'ok'
            out[url] = {'s': st, 't': title[:200]}
        except Exception as e:  # сеть/Google недоступны — оставляем прежний результат проверки
            out[url] = dict(prev.get(url) or {'s': 'err', 't': ''})
            print('  форма не открылась:', url, '-', str(e)[:100])
        time.sleep(0.3)
    return out


# ---------------------------------------------------------------- 2а. uSTEM Connect
CONNECT_LIST = 'https://%s/events?open=1' % CONNECT_HOST   # только события с открытой регистрацией
CONNECT_ALL = 'https://%s/events' % CONNECT_HOST           # все события, в том числе с закрытой регистрацией
CONNECT_LINK = re.compile(r'href="(?:https://%s)?/events/([a-z0-9][a-z0-9-]{2,120})/?"' % re.escape(CONNECT_HOST))
CONNECT_PAGE = re.compile(r'href="(?:https://%s)?/events/?\?(?:[^"]*&(?:amp;)?)?page=(\d{1,2})[^"]*"' % re.escape(CONNECT_HOST))


def parse_connect(page):
    """HTML списка событий Connect -> (адреса страниц событий, номера других страниц списка)."""
    slugs = []
    for m in CONNECT_LINK.finditer(page):
        if m.group(1) not in slugs:
            slugs.append(m.group(1))
    return slugs, {int(n) for n in CONNECT_PAGE.findall(page)}


def fetch_connect_list(url):
    """Адреса страниц всех событий одного списка Connect (со всех его страниц); [] — если список пуст."""
    slugs, seen, todo = [], {1}, [1]
    sep = '&' if '?' in url else '?'
    while todo and len(seen) <= 10:
        n = todo.pop(0)
        _, body = get(url + ('%spage=%d' % (sep, n) if n > 1 else ''), 20, cap=5 * 1024 * 1024)
        found, pages = parse_connect(body.decode('utf-8', 'replace'))
        slugs += [x for x in found if x not in slugs]
        for p in sorted(pages - seen):
            seen.add(p)
            todo.append(p)
    return sorted('https://%s/events/%s' % (CONNECT_HOST, x) for x in slugs)


def fetch_connect(prev, prev_all):
    """События uSTEM Connect -> (с открытой регистрацией, все). На сайте кнопка «Регистрация в uSTEM Connect» ведёт прямо
    на страницу своего турнира (сопоставление по названию — app.js → connMatch): регистрация открыта — с пометкой «открыта»,
    закрылась — без неё, но кнопка всё равно ведёт на страницу; турнира в Connect нет совсем — кнопка неактивна.
    Платформа недоступна или общий список пустой — остаётся прежний результат."""
    all_ev = fetch_connect_list(CONNECT_ALL)
    if not all_ev:
        print('   в списке Connect не найдено ни одного события — оставлены прежние ссылки')
        return prev, prev_all
    open_ev = fetch_connect_list(CONNECT_LIST)
    return open_ev, sorted(set(all_ev) | set(open_ev))


# ---------------------------------------------------------------- 3. Telegram
def strip_tags(s):
    s = re.sub(r'<br\s*/?>', '\n', s)
    s = re.sub(r'<[^>]+>', '', s)
    return html.unescape(s)


def parse_telegram(page, channel):
    posts = []
    for blk in re.split(r'(?=<div class="tgme_widget_message_wrap)', page)[1:]:
        m = re.search(r'data-post="%s/(\d+)"' % re.escape(channel), blk, re.I)
        if not m:
            continue
        pid = int(m.group(1))
        tm = re.search(r'<time[^>]+datetime="([^"]+)"', blk)
        tx = re.search(r'<div class="tgme_widget_message_text js-message_text"[^>]*>(.*?)</div>', blk, re.S)
        im = re.search(r'tgme_widget_message_(?:photo_wrap|video_thumb)[^>]*background-image:url\(\'([^\']+)\'\)', blk)
        posts.append({'id': pid, 'date': tm.group(1) if tm else '', 'text': strip_tags(tx.group(1)).strip() if tx else '',
                      'img_src': im.group(1) if im else ''})
    return posts


def shorten(s, n):
    s = re.sub(r'\s+', ' ', s).strip()
    if len(s) <= n:
        return s
    cut = s[:n].rsplit(' ', 1)[0].rstrip(' ,.;:—-')
    return cut + '…'


def news_item(p, channel, tags):
    text = p['text']
    text = re.sub(r'(?<!\w)#[\wәіңғүұқөһ]+', '', text)     # хэштеги на сайте не нужны
    lines = [l.strip() for l in text.split('\n') if l.strip()]
    if not lines:
        return None
    return {'id': p['id'], 'url': 'https://t.me/%s/%d' % (channel, p['id']), 'date': p['date'],
            'title': shorten(lines[0], 110), 'text': shorten(' '.join(lines[1:]), 220), 'img_src': p['img_src']}


def fetch_telegram(channel, tags, prev_items, max_items):
    found, before, pages = {}, None, 0
    while pages < 6:
        url = 'https://t.me/s/%s' % channel + ('?before=%d' % before if before else '')
        _, body = get(url, 30)
        posts = parse_telegram(body.decode('utf-8', 'replace'), channel)
        pages += 1
        if not posts:
            if pages == 1 and prev_items:
                # страница открылась, но разметка t.me не распозналась — старые новости не трогаем
                raise RuntimeError('страница t.me не разобралась — новости оставлены как были')
            break
        for p in posts:
            if any(re.search(re.escape(t) + r'(?![\wәіңғүұқөһ])', p['text'], re.I) for t in tags):
                it = news_item(p, channel, tags)
                if it:
                    found[it['id']] = it
        before = min(p['id'] for p in posts)
        if before <= 1:
            break
    items = {x['id']: x for x in prev_items if x.get('id')}
    items.update(found)          # посты, у которых хэштег убрали, остаются только если они старше просмотренных страниц
    seen_min = before or 0
    for pid in list(items):
        if pid >= seen_min and pid not in found:
            del items[pid]       # хэштег убрали или пост удалён — убираем с сайта
    out = sorted(items.values(), key=lambda x: x['date'], reverse=True)
    if max_items > 0:            # 0 в tools/config.json = без ограничения
        out = out[:max_items]
    os.makedirs(NEWS_IMG, exist_ok=True)
    keep = set()
    for it in out:
        fn = 'tg-%d.jpg' % it['id']
        path = os.path.join(NEWS_IMG, fn)
        src = it.pop('img_src', '') or ''
        if src.startswith('//'):
            src = 'https:' + src
        if src and not os.path.exists(path):
            try:
                _, img = get(src, 30, cap=5 * 1024 * 1024)
                if not (img.startswith(b'\xff\xd8\xff') or img.startswith(b'\x89PNG')):
                    raise RuntimeError('это не картинка')
                open(path, 'wb').write(img)
            except Exception as e:
                print('  фото не скачалось:', src[:80], str(e)[:80])
        if os.path.exists(path):
            it['img'] = 'assets/img/news/' + fn
            keep.add(fn)
        else:
            it.pop('img', None)
    for fn in os.listdir(NEWS_IMG):
        if fn.startswith('tg-') and fn not in keep:
            os.remove(os.path.join(NEWS_IMG, fn))
    return out


# ---------------------------------------------------------------- 4. вкладки «Сайт» и «Новости» -> site.js
# Вкладка «Сайт»: колонка A — имя параметра (по-русски или по-английски, регистр не важен), B/C/D — значение ru/kk/en;
# для ссылок и одиночных значений берётся только колонка B (у «Правила» C — казахская версия).
SITE_LABELS = {'season': 'Сезон', 'theme': 'Тема', 'logo': 'Логотип', 'video': 'Видео', 'rules': 'Правила', 'letter': 'Письмо',
               'schedule': 'Расписание', 'fll': 'FLL', 'ftc': 'FTC', 'frc': 'FRC'}
SITE_PARAMS = {k: k for k in SITE_LABELS}                       # английское имя (season, theme, logo, …)
SITE_PARAMS.update({v.lower(): k for k, v in SITE_LABELS.items()})   # русское имя (сезон, тема, логотип, …)
HEADER_WORDS = ('параметр', 'parameter', 'название', 'настройка', 'ключ', 'key')    # строка-шапка, если её сделали — не ругаемся
SITE_LIMITS = {'season': 60, 'theme': 300, 'fll': 40, 'ftc': 40, 'frc': 40}      # предел длины текстов, знаков


def site_text(c, key):
    s = shorten(txt(c), SITE_LIMITS[key])
    if s.endswith('…'):
        print('  вкладка «Сайт»: «%s» длиннее %d знаков — обрезано' % (SITE_LABELS[key], SITE_LIMITS[key]))
    return s


def site_link(c, key, hosts=LINK_HOSTS):
    u = cell_url(c)
    if u and not link_ok(u, hosts):
        print('  вкладка «Сайт»: ссылка «%s» пропущена — нужен https-адрес на %s: %s' % (key, ', '.join(hosts), u[:80]))
        return ''
    return u


def save_season_logo(url):
    """Логотип сезона -> assets/img/season.<ext> (один файл, прежний season.* удаляется). Возвращает путь для сайта или None.
    Если картинка не скачалась, прежний файл (если был) остаётся."""
    old = sorted(f for f in os.listdir(IMG) if re.match(r'^season\.(png|jpg|webp)$', f))
    if not url:
        for f in old:
            os.remove(os.path.join(IMG, f))
        return None
    try:
        data, ext = fetch_image(url)
    except Exception as e:
        print('  логотип не скачался:', url[:80], str(e)[:100], '— прежний оставлен' if old else '')
        return 'assets/img/' + old[0] if old else None
    fn = 'season.' + ext
    path = os.path.join(IMG, fn)
    if not (os.path.exists(path) and open(path, 'rb').read() == data):
        open(path + '.tmp', 'wb').write(data)
        os.replace(path + '.tmp', path)
    for f in old:
        if f != fn:
            os.remove(os.path.join(IMG, f))
    return 'assets/img/' + fn


def parse_site_tab(rows):
    """Строки вкладки «Сайт» -> раздел season для site.js."""
    vals = {}
    for r in rows:
        name = txt(r[0]).strip() if r else ''
        key = SITE_PARAMS.get(name.lower().rstrip(':').strip())
        if not name:
            continue
        if not key:
            if name.lower() not in HEADER_WORDS:
                print('  вкладка «Сайт»: неизвестный параметр «%s» пропущен (нужны: %s)' % (name[:40], ', '.join(SITE_LABELS.values())))
            continue
        vals[key] = [r[i] if i < len(r) else '' for i in (1, 2, 3)]    # B, C, D
    g = lambda k: vals.get(k, ['', '', ''])
    theme = {l: site_text(c, 'theme') for l, c in zip(('ru', 'kk', 'en'), g('theme'))}
    theme['ru'] = theme['ru'] or theme['kk'] or theme['en']
    theme = {l: theme[l] or theme['ru'] for l in theme} if theme['ru'] else None
    video = cell_url(g('video')[0])
    if video and not link_ok(video, VIDEO_HOSTS):
        print('  вкладка «Сайт»: «Видео» пропущено — нужна ссылка на youtube.com или youtu.be:', video[:80])
        video = ''
    sub = {k: (site_text(g(k)[0], k) or None) for k in ('fll', 'ftc', 'frc')}
    return {'name': site_text(g('season')[0], 'season') or None,
            'theme': theme,
            'logo': save_season_logo(cell_url(g('logo')[0])),
            'video': video or None,
            'rules': {'ru': site_link(g('rules')[0], 'Правила'), 'kk': site_link(g('rules')[1], 'Правила (қаз)')},
            'letter': site_link(g('letter')[0], 'Письмо') or None,
            'schedule': site_link(g('schedule')[0], 'Расписание') or 'https://docs.google.com/spreadsheets/d/%s/edit' % CFG['sheet_id'],
            'sub': sub}


# Вкладка «Новости»: колонки узнаются по шапке. Язык колонки — по пометке (ru)/(kk)/(en), рус/қаз/eng, иначе по языку самого
# слова в шапке (Заголовок / Тақырып / Title), иначе по порядку: первая — ru, вторая — kk, третья — en.
TITLE_RX = r'заголов|тақырып|title|headline|атауы|название'
TEXT_RX = r'текст|мәтін|text|описание|сипаттама'
LANG_MARK = (('ru', r'\bru\b|рус'), ('kk', r'\bk[kz]\b|қаз|каз'), ('en', r'\ben\b|eng|англ|ағыл'))
LANG_WORD = (('ru', r'заголов|текст|название|описание'), ('kk', r'тақырып|мәтін|атауы|сипаттама'), ('en', r'title|headline|text'))


def lang_cols(head, rx):
    """{'ru': i, 'kk': j, 'en': k} — номера колонок с заголовком по rx (языка может не хватать)."""
    found, free = {}, []
    for i, h in enumerate(head):
        h = txt(h).lower()
        if not re.search(rx, h):
            continue
        lang = next((l for l, m in LANG_MARK if re.search(m, h)), None) or next((l for l, m in LANG_WORD if re.search(m, h)), None)
        if lang and lang not in found:
            found[lang] = i
        else:
            free.append(i)
    for l in ('ru', 'kk', 'en'):
        if l not in found and free:
            found[l] = free.pop(0)
    return found


def find_col(head, rx):
    return next((i for i, h in enumerate(head) if re.search(rx, txt(h).lower())), -1)


def iso_date(c):
    """'15.10.2026', '15/10/2026', '2026-10-15' (и ячейка-дата) -> '2026-10-15'; иначе ''."""
    s = txt(c).strip()
    m = re.match(r'^(\d{1,2})[./-](\d{1,2})[./-](\d{4})', s)
    if m:
        d, mo, y = map(int, m.groups())
    else:
        m = re.match(r'^(\d{4})-(\d{1,2})-(\d{1,2})', s)
        if not m:
            return ''
        y, mo, d = map(int, m.groups())
    try:
        return datetime.date(y, mo, d).isoformat()
    except ValueError:
        return ''


def news_image(url):
    """Фото новости -> assets/img/news/sh-<hash>.<ext> (имя по адресу, повторно не скачивается). Путь для сайта или None."""
    url = (url or '').strip()
    if not url:
        return None
    stem = 'sh-' + hashlib.sha1(url.encode('utf-8')).hexdigest()[:10]
    old = [f for f in os.listdir(NEWS_IMG) if f.startswith(stem + '.')]
    if old:
        return 'assets/img/news/' + old[0]
    try:
        data, ext = fetch_image(url)
    except Exception as e:
        print('  фото новости не скачалось:', url[:80], str(e)[:100])
        return None
    fn = stem + '.' + ext
    open(os.path.join(NEWS_IMG, fn), 'wb').write(data)
    return 'assets/img/news/' + fn


def parse_news_tab(rows, max_items=12):
    """Строки вкладки «Новости» -> список новостей для site.js (новые сверху, не больше max_items)."""
    hi = next((i for i, r in enumerate(rows[:5]) if sum(1 for c in r if txt(c).strip()) >= 2), -1)   # шапка — первая строка с 2+ ячейками
    if hi < 0 or not lang_cols(rows[hi], TITLE_RX):
        raise RuntimeError('в шапке вкладки «Новости» нет колонки «Заголовок»')
    head = rows[hi]
    tc, xc = lang_cols(head, TITLE_RX), lang_cols(head, TEXT_RX)
    dc, uc, pc = date_col(head), find_col(head, r'ссылк|сілтеме|link|url'), find_col(head, r'фото|сурет|photo|image|картин')
    cell = lambda r, i: r[i] if 0 <= i < len(r) else ''
    os.makedirs(NEWS_IMG, exist_ok=True)
    items, keep = [], set()
    for r in rows[hi + 1:]:
        title = {l: shorten(txt(cell(r, tc.get(l, -1))), 200) for l in ('ru', 'kk', 'en')}
        title['ru'] = title['ru'] or title['kk'] or title['en']
        if not title['ru']:
            continue                                      # строка без заголовка — не новость
        text = {l: shorten(txt(cell(r, xc.get(l, -1))), 1000) for l in ('ru', 'kk', 'en')}
        text['ru'] = text['ru'] or text['kk'] or text['en']
        url = cell_url(cell(r, uc))
        if url and not https_ok(url):
            print('  вкладка «Новости»: ссылка пропущена (нужен https-адрес):', url[:80])
            url = ''
        items.append({'date': iso_date(cell(r, dc)), 'title': {l: title[l] or title['ru'] for l in title},
                      'text': {l: text[l] or text['ru'] for l in text}, 'url': url, 'img': cell_url(cell(r, pc))})
    items.sort(key=lambda x: x['date'], reverse=True)
    items = items[:max_items]
    for it in items:
        it['img'] = news_image(it['img'])
        if it['img']:
            keep.add(os.path.basename(it['img']))
    for fn in os.listdir(NEWS_IMG):
        if fn.startswith('sh-') and fn not in keep:
            os.remove(os.path.join(NEWS_IMG, fn))      # фото новостей, которых больше нет в таблице
    return items


def update_site(site_rows, news_rows, n_events, stamp):
    """assets/data/site.js = {updated, season|null, news, stats}. Сломанная вкладка — прежний раздел остаётся. Возвращает текст ошибки или ''."""
    prev = read_js('site.js', 'FK_SITE') or {}
    err = []
    season = prev.get('season')
    if site_rows is None:
        season = None
        print('   вкладки «Сайт» нет — сайт показывает встроенные тексты сезона')
    else:
        try:
            season = parse_site_tab(site_rows)
            print('   сезон: %s; тема: %s; логотип: %s; видео: %s' % (season['name'] or '—', 'да' if season['theme'] else '—', season['logo'] or '—', 'да' if season['video'] else '—'))
        except Exception as e:
            err.append('вкладка «Сайт»: %s' % str(e)[:100])
            print('   ОШИБКА вкладки «Сайт»:', str(e)[:150], '— прежние настройки оставлены')
    news = prev.get('news') or []
    if news_rows is None:
        news = []
    else:
        try:
            news = parse_news_tab(news_rows)
            print('   новостей из таблицы: %d' % len(news))
        except Exception as e:
            err.append('вкладка «Новости»: %s' % str(e)[:100])
            print('   ОШИБКА вкладки «Новости»:', str(e)[:150], '— прежние новости оставлены')
    site = {'updated': stamp, 'season': season, 'news': news, 'stats': {'events': n_events}}
    changed = {k: v for k, v in site.items() if k != 'updated'} != {k: v for k, v in prev.items() if k != 'updated'}
    if changed or (prev.get('updated') or '')[:7] != stamp[:7]:
        write_js('site.js', 'FK_SITE', site, 'Настройки сезона и новости из вкладок «Сайт» и «Новости» таблицы расписания. Обновляется автоматически (tools/update_data.py), вручную не редактировать.')
        print('   site.js обновлён' if changed else '   site.js без изменений (ежемесячная отметка)')
    else:
        print('   site.js без изменений')
    return '; '.join(err)


# ---------------------------------------------------------------- main
def main():
    os.makedirs(DATA, exist_ok=True)
    stamp = now_iso()
    old_status = read_js('status.js', 'FK_STATUS') or {}
    status = dict(old_status)
    prev = read_js('sheet.js', 'FK_SHEET') or {}
    rc = 0

    print('1. Таблица расписания…')
    try:
        tabs = fetch_sheet(CFG['sheet_id'])
        extra = {special_tab(t): tabs.pop(t) for t in list(tabs) if special_tab(t)}   # «Сайт» и «Новости» — не события
        same_sheet = (prev.get('source') or '').endswith('/%s/edit' % CFG['sheet_id'])   # новая таблица (сменили sheet_id) — защита от «случайного удаления» не нужна
        n = validate(tabs, prev.get('tabs') if same_sheet else None)
        print('   вкладок: %d, событий: %d' % (len(tabs), n))
        print('2. Формы регистрации…')
        quick = os.environ.get('FK_QUICK') == '1' and same_sheet and tabs == prev.get('tabs') and prev.get('forms') is not None
        if quick:   # быстрая проверка (каждые 5 минут): таблица не менялась — формы заново не открываем, их проверяет часовой запуск
            forms = prev['forms']
            print('   таблица не менялась — формы проверяются в часовом запуске')
        else:
            forms = check_forms(tabs, prev.get('forms') or {})
        bad = {u: f for u, f in forms.items() if f['s'] in ('closed', 'mismatch')}
        print('   форм: %d, скрыто: %d' % (len(forms), len(bad)))
        for u, f in bad.items():
            print('   [%s] %s — %s' % ('закрыта' if f['s'] == 'closed' else 'другое событие', u, f.get('t', '')))
        print('2а. uSTEM Connect…')
        try:
            connect, connect_all = fetch_connect(prev.get('connect') or [], prev.get('connect_all') or [])
            print('   событий в Connect: %d, с открытой регистрацией: %d' % (len(connect_all), len(connect)))
        except Exception as e:  # платформа недоступна — прежние ссылки остаются
            connect, connect_all = prev.get('connect') or [], prev.get('connect_all') or []
            print('   uSTEM Connect недоступен:', str(e)[:100])
        # 'updated' — то, что сайт показывает как «Обновлено»: меняется только при настоящем изменении таблицы, статуса формы
        # регистрации или списка событий Connect. (Раньше отметка ещё и переписывалась раз в месяц; репозиторий и без неё
        # не «засыпает» — status.js коммитится каждый день.)
        changed = (tabs != prev.get('tabs') or forms != prev.get('forms') or connect != (prev.get('connect') or [])
                   or connect_all != (prev.get('connect_all') or []))
        if changed:
            write_js('sheet.js', 'FK_SHEET', {'updated': stamp, 'source': 'https://docs.google.com/spreadsheets/d/%s/edit' % CFG['sheet_id'],
                                             'tabs': tabs, 'forms': forms, 'connect': connect, 'connect_all': connect_all},
                     'Снимок официальной таблицы расписания. Обновляется автоматически (tools/update_data.py), вручную не редактировать.')
            print('   sheet.js обновлён')
        else:
            print('   изменений нет')
        status['checked'] = stamp
        status['sheet'] = 'ok'
        print('3. Вкладки «Сайт» и «Новости»…')
        err = update_site(extra.get('site'), extra.get('news'), n, stamp)
        status['site'] = 'error: %s' % err[:150] if err else 'ok'
        rc = rc or (1 if err else 0)
    except Exception as e:
        print('   ОШИБКА таблицы:', e)
        status['sheet'] = 'error: %s' % str(e)[:150]
        rc = 1

    print('4. Telegram…')
    try:
        prev_news = read_js('news.js', 'FK_TG_NEWS') or []
        news = fetch_telegram(CFG['telegram_channel'], CFG.get('telegram_tags') or ['#сайт'], prev_news, int(CFG.get('telegram_max_news', 6)))
        if news != prev_news:
            write_js('news.js', 'FK_TG_NEWS', news, 'Посты Telegram-канала с хэштегом для сайта. Обновляется автоматически (tools/update_data.py).')
        print('   новостей с хэштегом: %d' % len(news))
        status['telegram'] = 'ok'
    except Exception as e:
        print('   Telegram недоступен:', e)
        status['telegram'] = 'error: %s' % str(e)[:150]

    # Сайт показывает только дату проверки (по времени Алматы), поэтому status.js переписывается, когда меняется эта дата
    # или результат проверки. Иначе каждый часовой запуск делал бы коммит только ради новой минуты в отметке.
    if status != old_status and (almaty_day(status.get('checked')) != almaty_day(old_status.get('checked'))
                                 or {k: v for k, v in status.items() if k != 'checked'} != {k: v for k, v in old_status.items() if k != 'checked'}):
        write_js('status.js', 'FK_STATUS', status, 'Время последней успешной проверки данных (tools/update_data.py).')
    return rc


def almaty_day(iso):
    """'2026-10-21T20:15:00Z' -> '2026-10-22' (дата в Алматы, UTC+5); пусто или не разобралось -> ''."""
    try:
        t = datetime.datetime.strptime(iso or '', '%Y-%m-%dT%H:%M:%SZ')
    except ValueError:
        return ''
    return (t + datetime.timedelta(hours=5)).date().isoformat()


if __name__ == '__main__':
    sys.exit(main())
